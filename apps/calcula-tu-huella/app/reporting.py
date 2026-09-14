from __future__ import annotations

import hashlib
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from .analytics import full_analysis
from .methodology_closure import closure_summary
from .storage import storage

from .database import (
    ActivityData,
    EmissionCalculation,
    EmissionFactorVersion,
    EmissionSource,
    Inventory,
    ReportArtifact,
    SupplierCampaign,
    SupplierDataRequest,
    SupplierResponse,
    INSTANCE_DIR,
)

REPORTS_DIR = INSTANCE_DIR / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

NAVY = colors.HexColor("#0F2D4D")
GREEN = colors.HexColor("#2E7D5B")
BLUE = colors.HexColor("#2D7DBD")
LIGHT = colors.HexColor("#F4F7F5")
GRID = colors.HexColor("#D9E3E0")
TEXT = colors.HexColor("#17232B")


def _money_cop(value: float) -> str:
    return "$" + f"{value:,.0f}".replace(",", ".")


def _number(value: float, decimals: int = 1) -> str:
    return f"{value:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _report_path(inventory: Inventory, suffix: str, label: str) -> Path:
    folder = REPORTS_DIR / f"org_{inventory.organization_id}" / f"inventory_{inventory.id}"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S_%f")
    safe_label = label.lower().replace(" ", "_").replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u")
    safe_name = f"{safe_label}_{inventory.start_date.year}_{inventory.id}_{stamp}.{suffix}"
    return folder / safe_name


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Brand", fontName="Helvetica-Bold", fontSize=18, textColor=NAVY, leading=22))
    styles.add(ParagraphStyle(name="Subtitle", fontName="Helvetica", fontSize=10, textColor=GREEN, leading=14))
    styles.add(ParagraphStyle(name="Section", fontName="Helvetica-Bold", fontSize=13, textColor=NAVY, spaceBefore=8, spaceAfter=6))
    styles.add(ParagraphStyle(name="BodySmall", fontName="Helvetica", fontSize=8.5, textColor=TEXT, leading=12))
    styles.add(ParagraphStyle(name="CenterSmall", fontName="Helvetica", fontSize=8, textColor=TEXT, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="Metric", fontName="Helvetica-Bold", fontSize=16, textColor=GREEN, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="MetricLabel", fontName="Helvetica", fontSize=7.5, textColor=NAVY, alignment=TA_CENTER))
    return styles


def _header_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(GRID)
    canvas.line(18 * mm, 14 * mm, 192 * mm, 14 * mm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#61727B"))
    canvas.drawString(18 * mm, 9 * mm, "Calcula tu Huella - documento generado desde la plataforma")
    canvas.drawRightString(192 * mm, 9 * mm, f"Pagina {doc.page}")
    canvas.restoreState()


def _table(data, widths=None, header=True, font_size=8):
    header_style = ParagraphStyle("TableHeader", fontName="Helvetica-Bold", fontSize=font_size, leading=font_size + 2, textColor=colors.white)
    body_style = ParagraphStyle("TableBody", fontName="Helvetica", fontSize=font_size, leading=font_size + 2, textColor=TEXT)
    prepared = []
    for row_index, row in enumerate(data):
        style = header_style if header and row_index == 0 else body_style
        prepared.append([cell if isinstance(cell, Paragraph) else Paragraph(str(cell), style) for cell in row])
    table = Table(prepared, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    commands = [
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("TEXTCOLOR", (0, 0), (-1, -1), TEXT),
        ("GRID", (0, 0), (-1, -1), 0.35, GRID),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        commands.extend([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ])
    for row in range(1 if header else 0, len(data)):
        if row % 2 == 0:
            commands.append(("BACKGROUND", (0, row), (-1, row), LIGHT))
    table.setStyle(TableStyle(commands))
    return table


def generate_executive_pdf(session: Session, inventory: Inventory, output: Path) -> None:
    analysis = full_analysis(session, inventory)
    closure = closure_summary(session, inventory)
    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(output), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=20 * mm,
        title=f"Informe ejecutivo {inventory.name}", author="Calcula tu Huella",
    )
    story = [
        Paragraph("CALCULA TU HUELLA", styles["Brand"]),
        Paragraph("Informe ejecutivo de huella de carbono", styles["Subtitle"]),
        Spacer(1, 8 * mm),
        Paragraph(inventory.organization.name, styles["Title"]),
        Paragraph(f"Periodo: {inventory.start_date:%d/%m/%Y} - {inventory.end_date:%d/%m/%Y}", styles["BodySmall"]),
        Paragraph(f"Metodologia: {inventory.methodology} | GWP: {inventory.gwp_version}", styles["BodySmall"]),
        Spacer(1, 8 * mm),
    ]
    history = analysis["history"]
    total_change = history["total_change"]
    metrics_data = [
        [Paragraph("EMISIONES TOTALES", styles["MetricLabel"]), Paragraph("INTENSIDAD PRODUCTIVA", styles["MetricLabel"]), Paragraph("CALIDAD DEL DATO", styles["MetricLabel"]), Paragraph("VARIACION ANUAL", styles["MetricLabel"])],
        [Paragraph(f"{_number(analysis['total'])}<br/><font size=8>tCO2e</font>", styles["Metric"]),
         Paragraph(f"{_number(analysis['intensity_production'] or 0, 4)}<br/><font size=8>tCO2e/t</font>", styles["Metric"]),
         Paragraph(f"{analysis['quality']['score']}%", styles["Metric"]),
         Paragraph("N/D" if total_change is None else f"{total_change:+.1f}%", styles["Metric"])],
    ]
    metric_table = Table(metrics_data, colWidths=[43 * mm] * 4)
    metric_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), LIGHT),
        ("BOX", (0, 0), (-1, -1), 0.5, GRID),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, GRID),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([metric_table, Spacer(1, 7 * mm)])
    balance = closure["balance"]
    uncertainty = closure["uncertainty"]
    story.append(Paragraph("Balance metodológico", styles["Section"]))
    methodology_rows = [
        ["Partida", "tCO2e", "Tratamiento"],
        ["Emisiones brutas", _number(balance["gross_emissions"], 3), "Inventario corporativo"],
        ["CO2 biogénico", _number(balance["biogenic_memo"], 3), "Partida informativa"],
        ["Remociones", _number(balance["removals"], 3), "Separadas del bruto"],
        ["Emisiones evitadas", _number(balance["avoided_emissions"], 3), "Fuera del inventario físico"],
        ["Compensaciones", _number(balance["offsets"], 3), "Fuera del inventario bruto"],
    ]
    story.extend([_table(methodology_rows, [65 * mm, 38 * mm, 62 * mm], font_size=7.5)])
    story.append(Paragraph(
        f"Incertidumbre combinada sobre emisiones cubiertas: {uncertainty['combined_percentage']:.2f}%. "
        f"Rango cubierto: {_number(uncertainty['lower_tco2e'], 3)} a {_number(uncertainty['upper_tco2e'], 3)} tCO2e "
        f"({uncertainty['emission_coverage_percentage']:.1f}% del total bruto). "
        f"Preparación metodológica: {closure['readiness_score']}%.",
        styles["BodySmall"],
    ))
    story.extend([Spacer(1, 7 * mm), Paragraph("Resultados por alcance", styles["Section"])])
    scopes = analysis["scopes"]
    scope_rows = [["Alcance", "tCO2e", "Participacion"]]
    for scope in (1, 2, 3):
        share = scopes[scope] / analysis["total"] * 100 if analysis["total"] else 0
        scope_rows.append([f"Alcance {scope}", _number(scopes[scope]), f"{share:.1f}%"])
    scope_rows.append(["Total", _number(analysis["total"]), "100,0%"])
    story.extend([_table(scope_rows, [70 * mm, 45 * mm, 45 * mm]), Spacer(1, 7 * mm)])

    story.append(Paragraph("Principales fuentes", styles["Section"]))
    source_rows = [["Fuente", "Alcance", "Sede", "tCO2e", "%"]]
    for row in analysis["sources_summary"][:8]:
        source_rows.append([row["name"], row["scope"], row["facility"], _number(row["emissions"]), f"{row['share']:.1f}%"])
    story.extend([_table(source_rows, [42 * mm, 18 * mm, 45 * mm, 30 * mm, 20 * mm]), Spacer(1, 7 * mm)])

    story.append(Paragraph("Lectura ejecutiva", styles["Section"]))
    top_sources = analysis["sources_summary"][:2]
    observations = []
    if top_sources:
        observations.append(f"Las dos fuentes principales concentran {sum(row['share'] for row in top_sources):.1f}% de las emisiones.")
    if total_change is not None:
        direction = "aumentaron" if total_change > 0 else "disminuyeron"
        observations.append(f"Las emisiones absolutas {direction} {abs(total_change):.1f}% frente al inventario anterior.")
    intensity_change = history["intensity_change"]
    if intensity_change is not None:
        direction = "aumento" if intensity_change > 0 else "disminuyo"
        observations.append(f"La intensidad de carbono {direction} {abs(intensity_change):.1f}% frente al periodo anterior.")
    observations.append(f"La calidad consolidada de datos es {analysis['quality']['score']}%, con cobertura documental de {analysis['quality']['evidence_coverage']}%.")
    for item in observations:
        story.append(Paragraph(f"- {item}", styles["BodySmall"]))
    story.extend([Spacer(1, 6 * mm), Paragraph("Plan de reduccion", styles["Section"])])
    reduction = analysis["reduction"]
    reduction_rows = [["Accion", "Fuente", "Reduccion esperada", "Inversion", "Estado"]]
    for action in reduction["actions"]:
        reduction_rows.append([
            action.title,
            action.source.name if action.source else "Corporativo",
            f"{_number(action.expected_reduction)} tCO2e/anio",
            _money_cop(action.investment_cost),
            action.status,
        ])
    story.extend([_table(reduction_rows, [47 * mm, 30 * mm, 34 * mm, 32 * mm, 28 * mm], font_size=7.2), Spacer(1, 5 * mm)])
    story.append(Paragraph(
        "Nota: Este documento corresponde a un informe generado por la plataforma con base en la informacion registrada. "
        "La revision interna no equivale a verificacion independiente.", styles["BodySmall"]
    ))
    doc.build(story, onFirstPage=_header_footer, onLaterPages=_header_footer)


def generate_technical_pdf(session: Session, inventory: Inventory, output: Path) -> None:
    analysis = full_analysis(session, inventory)
    closure = closure_summary(session, inventory)
    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(output), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm,
        topMargin=18 * mm, bottomMargin=20 * mm,
        title=f"Informe tecnico {inventory.name}", author="Calcula tu Huella",
    )
    story = [
        Paragraph("CALCULA TU HUELLA", styles["Brand"]),
        Paragraph("Informe tecnico del inventario corporativo de GEI", styles["Subtitle"]),
        Spacer(1, 6 * mm),
        Paragraph(inventory.organization.name, styles["Title"]),
        Paragraph(f"Inventario: {inventory.name} | Version: {inventory.version}", styles["BodySmall"]),
        Paragraph(f"Periodo: {inventory.start_date:%d/%m/%Y} - {inventory.end_date:%d/%m/%Y}", styles["BodySmall"]),
        PageBreak(),
        Paragraph("1. Perfil y objetivo", styles["Section"]),
        Paragraph(
            f"La organizacion pertenece al sector {inventory.organization.sector}, con {inventory.organization.employees} empleados. "
            f"El objetivo declarado del inventario es: {inventory.objective}", styles["BodySmall"]
        ),
        Paragraph("2. Limites y metodologia", styles["Section"]),
        Paragraph(
            f"Metodologia: {inventory.methodology}. Version metodologica: {inventory.methodology_version}. "
            f"Consolidacion: {inventory.consolidation_approach}. GWP: {inventory.gwp_version}. "
            f"Umbral de materialidad: {inventory.materiality_threshold:.1f}%.", styles["BodySmall"]
        ),
        Paragraph("3. Fuentes incluidas", styles["Section"]),
    ]
    source_rows = [["Fuente", "Alcance", "Categoria", "Sede", "Materialidad", "tCO2e"]]
    for row in analysis["sources_summary"]:
        source = next(item for item in inventory.sources if item.id == row["id"])
        source_rows.append([row["name"], row["scope"], row["category"], row["facility"], source.materiality, _number(row["emissions"], 3)])
    story.extend([_table(source_rows, [30 * mm, 14 * mm, 38 * mm, 38 * mm, 25 * mm, 24 * mm], font_size=7), PageBreak()])

    story.append(Paragraph("4. Resultados consolidados", styles["Section"]))
    scope_rows = [["Alcance", "tCO2e"]] + [[f"Alcance {scope}", _number(value, 3)] for scope, value in analysis["scopes"].items()] + [["Total", _number(analysis["total"], 3)]]
    story.extend([_table(scope_rows, [80 * mm, 50 * mm]), Spacer(1, 6 * mm)])

    balance = closure["balance"]
    uncertainty = closure["uncertainty"]
    scope2 = closure["scope2"]
    story.append(Paragraph("4.1. Tratamiento contable y partidas separadas", styles["Section"]))
    accounting_rows = [
        ["Partida", "tCO2e", "Presentación"],
        ["Emisiones brutas", _number(balance["gross_emissions"], 6), "Total principal por alcances"],
        ["CO2 biogénico", _number(balance["biogenic_memo"], 6), "Memorando informativo"],
        ["Remociones", _number(balance["removals"], 6), "Separadas; no reducen el bruto"],
        ["Emisiones evitadas", _number(balance["avoided_emissions"], 6), "Fuera del inventario físico"],
        ["Compensaciones", _number(balance["offsets"], 6), "Fuera del inventario bruto"],
        ["Neto de referencia tras remociones", _number(balance["net_after_removals"], 6), "Indicador complementario"],
    ]
    story.extend([_table(accounting_rows, [58 * mm, 35 * mm, 77 * mm], font_size=7.2), Spacer(1, 5 * mm)])

    story.append(Paragraph("4.2. Alcance 2 e incertidumbre", styles["Section"]))
    scope2_rows = [
        ["Métrica", "Resultado"],
        ["Alcance 2 location-based", f"{_number(scope2['location_based'], 6)} tCO2e"],
        ["Alcance 2 market-based", f"{_number(scope2['market_based'], 6)} tCO2e"],
        ["Fuentes de alcance 2 sin clasificar", str(len(scope2['unclassified']))],
        ["Incertidumbre combinada", f"{uncertainty['combined_percentage']:.2f}%"],
        ["Rango inferior de emisiones cubiertas", f"{_number(uncertainty['lower_tco2e'], 6)} tCO2e"],
        ["Rango superior de emisiones cubiertas", f"{_number(uncertainty['upper_tco2e'], 6)} tCO2e"],
        ["Cobertura de emisiones brutas", f"{uncertainty['emission_coverage_percentage']}%"],
    ]
    story.extend([_table(scope2_rows, [90 * mm, 70 * mm]), Spacer(1, 5 * mm)])
    readiness_rows = [["Puerta metodológica", "Estado"]] + [[item["name"], item["status"]] for item in closure["readiness"]]
    story.extend([_table(readiness_rows, [105 * mm, 55 * mm], font_size=7.2), Spacer(1, 6 * mm)])

    story.append(Paragraph("5. Indicadores de intensidad", styles["Section"]))
    indicator_rows = [["Indicador", "Valor", "Unidad", "Intensidad"]]
    for name, metric in analysis["indicators"].items():
        if name == "Producción":
            intensity = analysis["intensity_production"]
            intensity_label = f"{_number(intensity or 0, 6)} tCO2e/{metric.unit}"
        elif name == "Empleados":
            intensity_label = f"{_number(analysis['intensity_employee'] or 0, 6)} tCO2e/persona"
        elif name == "Ingresos":
            intensity_label = f"{_number(analysis['intensity_revenue'] or 0, 6)} tCO2e/millon COP"
        else:
            intensity_label = "-"
        indicator_rows.append([name, _number(metric.value, 2), metric.unit, intensity_label])
    story.extend([_table(indicator_rows, [40 * mm, 35 * mm, 35 * mm, 60 * mm]), Spacer(1, 6 * mm)])

    quality = analysis["quality"]
    story.append(Paragraph("6. Calidad de los datos", styles["Section"]))
    quality_rows = [["Criterio", "Resultado"], ["Puntaje consolidado", f"{quality['score']}%"], ["Cobertura documental", f"{quality['evidence_coverage']}%"], ["Datos estimados", f"{quality['estimated_share']}%"]]
    for level, count in quality["counts"].items():
        quality_rows.append([f"Registros nivel {level}", str(count)])
    story.extend([_table(quality_rows, [85 * mm, 65 * mm]), Spacer(1, 6 * mm)])

    story.append(Paragraph("7. Cadena de valor y proveedores", styles["Section"]))
    supplier_responses = list(session.scalars(
        select(SupplierResponse)
        .join(SupplierDataRequest).join(SupplierCampaign)
        .where(SupplierCampaign.inventory_id == inventory.id)
        .options(selectinload(SupplierResponse.request).selectinload(SupplierDataRequest.supplier))
        .order_by(SupplierResponse.calculated_emissions_tco2e.desc())
    ))
    supplier_rows = [["Proveedor", "Producto/servicio", "Metodo", "Calidad", "Revision", "tCO2e"]]
    for response in supplier_responses:
        supplier_rows.append([response.request.supplier.name, response.request.product_service, response.method, response.quality_level, response.review_status, _number(response.calculated_emissions_tco2e, 3)])
    if len(supplier_rows) == 1:
        supplier_rows.append(["Sin respuestas", "-", "-", "-", "-", "0"] )
    story.extend([_table(supplier_rows, [38 * mm, 42 * mm, 34 * mm, 18 * mm, 24 * mm, 22 * mm], font_size=6.7), PageBreak()])

    story.append(Paragraph("8. Memoria resumida de calculos", styles["Section"]))
    calculations = list(
        session.scalars(
            select(EmissionCalculation)
            .join(ActivityData)
            .join(EmissionSource)
            .where(EmissionSource.inventory_id == inventory.id)
            .options(selectinload(EmissionCalculation.activity_data), selectinload(EmissionCalculation.factor_version).selectinload(EmissionFactorVersion.factor))
            .order_by(ActivityData.period_start, EmissionCalculation.id)
        )
    )
    calc_rows = [["Periodo", "Fuente", "Dato", "Factor", "Gas", "kg CO2e", "Estado"]]
    for calc in calculations:
        calc_rows.append([
            calc.activity_data.period_start.strftime("%Y-%m"),
            calc.activity_data.source.name,
            f"{_number(calc.original_value, 3)} {calc.original_unit}",
            calc.factor_version.factor.name,
            calc.gas_code,
            _number(calc.co2e_kg, 3),
            calc.status,
        ])
    story.extend([_table(calc_rows, [21 * mm, 26 * mm, 25 * mm, 47 * mm, 15 * mm, 22 * mm, 20 * mm], font_size=6.2), PageBreak()])

    story.append(Paragraph("9. Acciones de reduccion", styles["Section"]))
    reduction_rows = [["Accion", "Fuente", "Reduccion", "Inversion", "Ahorro anual", "Avance"]]
    for action in analysis["reduction"]["actions"]:
        reduction_rows.append([
            action.title,
            action.source.name if action.source else "Corporativo",
            f"{_number(action.expected_reduction)} tCO2e",
            _money_cop(action.investment_cost),
            _money_cop(action.annual_savings),
            f"{action.progress_percent}%",
        ])
    story.extend([_table(reduction_rows, [44 * mm, 29 * mm, 26 * mm, 32 * mm, 32 * mm, 18 * mm], font_size=6.8), Spacer(1, 6 * mm)])
    story.append(Paragraph("10. Declaracion tecnica", styles["Section"]))
    story.append(Paragraph(
        "El inventario fue elaborado a partir de la informacion registrada por la organizacion y los factores seleccionados. "
        "Los factores demostrativos deben sustituirse por fuentes oficiales o especificas antes de uso externo. "
        "La aprobacion dentro de la plataforma corresponde a control interno y no constituye verificacion independiente.", styles["BodySmall"]
    ))
    doc.build(story, onFirstPage=_header_footer, onLaterPages=_header_footer)


def generate_calculation_workbook(session: Session, inventory: Inventory, output: Path) -> None:
    analysis = full_analysis(session, inventory)
    closure = closure_summary(session, inventory)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["CALCULA TU HUELLA - MEMORIA DE CALCULO"])
    ws.append(["Organización", inventory.organization.name])
    ws.append(["Inventario", inventory.name])
    ws.append(["Periodo", f"{inventory.start_date:%Y-%m-%d} a {inventory.end_date:%Y-%m-%d}"])
    ws.append(["Metodología", inventory.methodology])
    ws.append(["GWP", inventory.gwp_version])
    ws.append([])
    ws.append(["Indicador", "Valor", "Unidad"])
    ws.append(["Emisiones totales", analysis["total"], "tCO2e"])
    for scope, value in analysis["scopes"].items():
        ws.append([f"Alcance {scope}", value, "tCO2e"])
    ws.append(["Calidad de datos", analysis["quality"]["score"], "%"])
    ws.append(["Intensidad productiva", analysis["intensity_production"] or 0, "tCO2e/t"])
    ws.append(["Incertidumbre combinada", closure["uncertainty"]["combined_percentage"], "%"])
    ws.append(["Rango inferior (emisiones cubiertas)", closure["uncertainty"]["lower_tco2e"], "tCO2e"])
    ws.append(["Rango superior (emisiones cubiertas)", closure["uncertainty"]["upper_tco2e"], "tCO2e"])
    ws.append(["Preparación metodológica", closure["readiness_score"], "%"])

    data_ws = wb.create_sheet("Datos de actividad")
    data_ws.append(["ID", "Fuente", "Sede", "Periodo inicial", "Periodo final", "Valor", "Unidad", "Origen", "Calidad", "Estimado", "Incertidumbre %", "Base incertidumbre", "Evidencia", "Estado", "Notas"])
    records = list(
        session.scalars(
            select(ActivityData)
            .join(EmissionSource)
            .where(EmissionSource.inventory_id == inventory.id)
            .options(selectinload(ActivityData.source).selectinload(EmissionSource.facility), selectinload(ActivityData.evidence))
            .order_by(ActivityData.period_start, ActivityData.id)
        )
    )
    for record in records:
        data_ws.append([
            record.id, record.source.name, record.source.facility.name if record.source.facility else "Corporativo",
            record.period_start, record.period_end, record.value, record.unit, record.data_origin, record.quality_level,
            "Sí" if record.is_estimated else "No", record.uncertainty_percentage, record.uncertainty_basis,
            record.evidence.name if record.evidence else "", record.status, record.notes,
        ])

    calc_ws = wb.create_sheet("Cálculos")
    calc_ws.append(["ID", "Dato ID", "Fuente", "Periodo", "Valor original", "Unidad", "Valor normalizado", "Unidad normalizada", "Factor", "Versión", "Gas", "Resultado gas kg", "GWP", "kg CO2e", "tCO2e", "Partida", "Incertidumbre %", "Límite inferior kg", "Límite superior kg", "Estado", "Alerta", "Fórmula"])
    calculations = list(
        session.scalars(
            select(EmissionCalculation)
            .join(ActivityData)
            .join(EmissionSource)
            .where(EmissionSource.inventory_id == inventory.id)
            .options(selectinload(EmissionCalculation.activity_data).selectinload(ActivityData.source), selectinload(EmissionCalculation.factor_version).selectinload(EmissionFactorVersion.factor))
            .order_by(ActivityData.period_start, EmissionCalculation.id)
        )
    )
    for calc in calculations:
        calc_ws.append([
            calc.id, calc.activity_data_id, calc.activity_data.source.name, calc.activity_data.period_start,
            calc.original_value, calc.original_unit, calc.normalized_value, calc.normalized_unit,
            calc.factor_version.factor.name, calc.factor_version.version, calc.gas_code, calc.gas_result_kg,
            calc.gwp_value, calc.co2e_kg, calc.co2e_kg / 1000, calc.reporting_bucket, calc.uncertainty_percentage,
            calc.lower_co2e_kg, calc.upper_co2e_kg, calc.status, calc.warning, calc.formula_snapshot,
        ])

    factors_ws = wb.create_sheet("Factores")
    factors_ws.append(["Fuente", "Factor", "Versión", "Valor", "Unidad entrada", "Unidad salida", "Gas", "Incertidumbre %", "Organización fuente", "Documento", "Año", "Estado", "Demostrativo"])
    seen = set()
    for source in inventory.sources:
        for assignment in source.factor_assignments:
            version = assignment.factor_version
            key = (source.id, version.id)
            if key in seen:
                continue
            seen.add(key)
            factors_ws.append([
                source.name, version.factor.name, version.version, version.value, version.input_unit, version.output_unit,
                version.gas.code, version.uncertainty_percentage, version.source_organization, version.source_document, version.publication_year,
                version.status, "Sí" if version.factor.is_demo else "No",
            ])

    method_ws = wb.create_sheet("Cierre metodológico")
    method_ws.append(["Componente", "Valor", "Unidad / estado"])
    balance = closure["balance"]
    method_ws.append(["Emisiones brutas", balance["gross_emissions"], "tCO2e"])
    method_ws.append(["CO2 biogénico informativo", balance["biogenic_memo"], "tCO2e"])
    method_ws.append(["Remociones", balance["removals"], "tCO2e"])
    method_ws.append(["Emisiones evitadas", balance["avoided_emissions"], "tCO2e"])
    method_ws.append(["Compensaciones", balance["offsets"], "tCO2e"])
    method_ws.append(["Alcance 2 location-based", closure["scope2"]["location_based"], "tCO2e"])
    method_ws.append(["Alcance 2 market-based", closure["scope2"]["market_based"], "tCO2e"])
    method_ws.append(["Fuentes alcance 2 sin clasificar", len(closure["scope2"]["unclassified"]), "cantidad"])
    method_ws.append(["Incertidumbre combinada (emisiones cubiertas)", closure["uncertainty"]["combined_percentage"], "%"])
    method_ws.append(["Cobertura de emisiones brutas", closure["uncertainty"]["emission_coverage_percentage"], "%"])
    method_ws.append(["Política metodológica", closure["policy"].get("status", "Borrador"), "estado"])
    method_ws.append([])
    method_ws.append(["Puerta metodológica", "Estado", "Acción requerida"])
    for item in closure["readiness"]:
        method_ws.append([item["name"], item["status"], item["detail"]])

    indicator_ws = wb.create_sheet("Indicadores")
    indicator_ws.append(["Tipo", "Periodo inicial", "Periodo final", "Sede", "Valor", "Unidad", "Fuente", "Estado", "Notas"])
    for indicator in inventory.indicators:
        indicator_ws.append([
            indicator.indicator_type, indicator.period_start, indicator.period_end,
            indicator.facility.name if indicator.facility else "Corporativo", indicator.value, indicator.unit,
            indicator.source_name, indicator.status, indicator.notes,
        ])

    supplier_ws = wb.create_sheet("Proveedores alcance 3")
    supplier_ws.append(["Campaña", "Proveedor", "NIT", "Producto/servicio", "Cantidad", "Unidad", "Gasto COP", "Método", "Factor", "Unidad factor", "tCO2e", "Metodología", "Límite", "Verificada", "Calidad", "Revisión", "SHA256 evidencia"])
    supplier_responses = list(session.scalars(
        select(SupplierResponse)
        .join(SupplierDataRequest).join(SupplierCampaign)
        .where(SupplierCampaign.inventory_id == inventory.id)
        .options(selectinload(SupplierResponse.request).selectinload(SupplierDataRequest.supplier), selectinload(SupplierResponse.request).selectinload(SupplierDataRequest.campaign))
        .order_by(SupplierResponse.id)
    ))
    for response in supplier_responses:
        req = response.request
        supplier_ws.append([req.campaign.name, req.supplier.name, req.supplier.tax_id, req.product_service, req.quantity, req.unit, req.spend_cop, response.method, response.emission_factor, response.factor_unit, response.calculated_emissions_tco2e, response.methodology, response.boundary, "Sí" if response.verified else "No", response.quality_level, response.review_status, response.evidence_sha256])

    reduction_ws = wb.create_sheet("Reducción")
    reduction_ws.append(["Acción", "Fuente", "Descripción", "Línea base tCO2e", "Reducción esperada", "Inversión COP", "Ahorro anual COP", "Prioridad", "Responsable", "Fecha objetivo", "Estado", "Avance %", "Reducción real", "Ahorro real"])
    for action in inventory.reduction_actions:
        reduction_ws.append([
            action.title, action.source.name if action.source else "Corporativo", action.description,
            action.baseline_emissions, action.expected_reduction, action.investment_cost, action.annual_savings,
            action.priority, action.responsible, action.target_date, action.status, action.progress_percent,
            action.actual_reduction, action.actual_savings,
        ])

    header_fill = PatternFill("solid", fgColor="0F2D4D")
    header_font = Font(color="FFFFFF", bold=True)
    title_fill = PatternFill("solid", fgColor="EAF3EF")
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Resumen" else "A8"
        header_row = 1 if sheet.title != "Resumen" else 8
        for cell in sheet[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        if sheet.title == "Resumen":
            sheet["A1"].fill = title_fill
            sheet["A1"].font = Font(bold=True, size=16, color="0F2D4D")
        for column_cells in sheet.columns:
            max_len = 0
            for cell in list(column_cells)[:200]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 10), 42)
    wb.save(output)


def create_report_artifact(session: Session, inventory: Inventory, report_type: str, generated_by: str) -> ReportArtifact:
    normalized = report_type.lower()
    if normalized == "ejecutivo":
        output = _report_path(inventory, "pdf", "informe_ejecutivo")
        generate_executive_pdf(session, inventory, output)
        label = "Informe ejecutivo"
    elif normalized == "tecnico":
        output = _report_path(inventory, "pdf", "informe_tecnico")
        generate_technical_pdf(session, inventory, output)
        label = "Informe técnico"
    elif normalized == "memoria":
        output = _report_path(inventory, "xlsx", "memoria_calculo")
        generate_calculation_workbook(session, inventory, output)
        label = "Memoria de cálculo"
    else:
        raise ValueError("Tipo de informe no soportado")
    content = output.read_bytes()
    storage_key = str(output.relative_to(INSTANCE_DIR))
    storage.put_bytes(storage_key, content, "application/pdf" if output.suffix.lower() == ".pdf" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    artifact = ReportArtifact(
        inventory_id=inventory.id,
        report_type=label,
        version=f"{inventory.version}-{datetime.now(UTC):%Y%m%d%H%M%S}",
        status="Generado",
        file_name=output.name,
        stored_name=storage_key,
        file_size=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        generated_by=generated_by,
    )
    session.add(artifact)
    session.flush()
    return artifact
